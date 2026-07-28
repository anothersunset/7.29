"""Solve, validate, and visualize the 50x50 puzzle from Attachment 2."""
from __future__ import annotations

import argparse
import ast
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

from src.analysis import analyze_solutions
from src.encoding import encode_line, parse_clue
from src.nonogram_dp import Matrix, solve_nonogram


def parse_problem4_text(path: Path) -> tuple[tuple[tuple[int, ...], ...], tuple[tuple[int, ...], ...]]:
    lines = [
        line.strip()
        for line in path.read_text(encoding="utf-8-sig").splitlines()
        if line.strip()
    ]
    size: int | None = None
    section: str | None = None
    rows: list[tuple[int, ...]] = []
    cols: list[tuple[int, ...]] = []
    for line in lines:
        lower = line.lower()
        if lower.startswith("size:"):
            dimensions = lower.split(":", 1)[1].split("*")
            if len(dimensions) != 2 or dimensions[0] != dimensions[1]:
                raise ValueError("size must describe a square matrix")
            size = int(dimensions[0])
        elif lower == "col:":
            section = "col"
        elif lower == "row:":
            section = "row"
        elif line.startswith("["):
            if section is None:
                raise ValueError("clue appears before Row/Col section")
            value = ast.literal_eval(line)
            if not isinstance(value, list):
                raise ValueError(f"invalid clue line: {line}")
            clue = parse_clue(value)
            (cols if section == "col" else rows).append(clue)
    if size is None or len(rows) != size or len(cols) != size:
        raise ValueError(
            f"expected a complete square puzzle, got size={size}, rows={len(rows)}, cols={len(cols)}"
        )
    return tuple(rows), tuple(cols)


def validate_solution(
    matrix: Matrix,
    rows: tuple[tuple[int, ...], ...],
    cols: tuple[tuple[int, ...], ...],
) -> bool:
    size = len(rows)
    return (
        tuple(encode_line(row) for row in matrix) == rows
        and tuple(
            encode_line(tuple(matrix[row_index][col_index] for row_index in range(size)))
            for col_index in range(size)
        )
        == cols
    )


def render_matrix(
    matrix: tuple[tuple[int | None, ...], ...],
    path: Path,
    scale: int = 20,
) -> None:
    size = len(matrix)
    image = Image.new("RGB", (size * scale, size * scale), "white")
    draw = ImageDraw.Draw(image)
    colors = {0: "white", 1: "black", None: "#ef4444"}
    for row_index, row in enumerate(matrix):
        for col_index, value in enumerate(row):
            draw.rectangle(
                (
                    col_index * scale,
                    row_index * scale,
                    (col_index + 1) * scale - 1,
                    (row_index + 1) * scale - 1,
                ),
                fill=colors[value],
            )
    image.save(path)


def write_workbook(
    path: Path,
    rows: tuple[tuple[int, ...], ...],
    cols: tuple[tuple[int, ...], ...],
    solutions: tuple[Matrix, ...],
) -> None:
    analysis = analyze_solutions(solutions)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary.append(["指标", "结果"])
    summary.append(["矩阵规模", f"{len(rows)}×{len(rows)}"])
    summary.append(["完整解数量", analysis.solution_count])
    summary.append(["是否唯一", "是" if analysis.is_unique else "否"])
    summary.append(["多解空格数", analysis.uncertain_count])
    summary.append(["所有解行列验证", "通过"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCEEFF")
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 22

    black = PatternFill("solid", fgColor="111827")
    red = PatternFill("solid", fgColor="EF4444")
    blue = PatternFill("solid", fgColor="DCEEFF")
    center = Alignment(horizontal="center", vertical="center")

    def add_matrix_sheet(title: str, matrix: tuple[tuple[int | None, ...], ...]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B2"
        sheet.cell(1, 1, "行\\列")
        for col_index, clue in enumerate(cols, 2):
            sheet.cell(1, col_index, ",".join(map(str, clue)))
        for row_index, clue in enumerate(rows, 2):
            sheet.cell(row_index, 1, ",".join(map(str, clue)))
        for row_index, row in enumerate(matrix, 2):
            for col_index, value in enumerate(row, 2):
                cell = sheet.cell(row_index, col_index, "?" if value is None else value)
                cell.alignment = center
                if value == 1:
                    cell.fill = black
                    cell.font = Font(color="FFFFFF")
                elif value is None:
                    cell.fill = red
                    cell.font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = blue
            cell.font = Font(bold=True)
            cell.alignment = center
        for row_index in range(2, len(rows) + 2):
            sheet.cell(row_index, 1).fill = blue
            sheet.cell(row_index, 1).font = Font(bold=True)
            sheet.cell(row_index, 1).alignment = center
        sheet.column_dimensions["A"].width = 24
        for col_index in range(2, len(cols) + 2):
            sheet.column_dimensions[get_column_letter(col_index)].width = 3
        for row_index in range(2, len(rows) + 2):
            sheet.row_dimensions[row_index].height = 18

    add_matrix_sheet("公共矩阵", analysis.consensus_matrix)
    add_matrix_sheet("代表解", solutions[0])
    ambiguous = workbook.create_sheet("多解空格")
    ambiguous.append(["序号", "行", "列"])
    for index, (row_index, col_index) in enumerate(analysis.uncertain_cells, 1):
        ambiguous.append([index, row_index, col_index])
    workbook.save(path)


def run(input_path: Path, output_dir: Path) -> dict[str, object]:
    rows, cols = parse_problem4_text(input_path)
    result = solve_nonogram(rows, cols)
    if not result.exhausted or not result.solutions:
        raise RuntimeError("Problem 4 search did not finish with a solution")
    validations = [validate_solution(matrix, rows, cols) for matrix in result.solutions]
    if not all(validations):
        raise RuntimeError("at least one reconstructed matrix failed validation")

    analysis = analyze_solutions(result.solutions)
    output_dir.mkdir(parents=True, exist_ok=True)
    render_matrix(result.solutions[0], output_dir / "problem4_representative.png")
    render_matrix(analysis.consensus_matrix, output_dir / "problem4_consensus.png")
    write_workbook(
        output_dir / "problem4_results.xlsx",
        rows,
        cols,
        result.solutions,
    )
    payload: dict[str, object] = {
        "input_file": input_path.name,
        "size": len(rows),
        "solution_count": analysis.solution_count,
        "is_unique": analysis.is_unique,
        "uncertain_count": analysis.uncertain_count,
        "uncertain_cells": [list(cell) for cell in analysis.uncertain_cells],
        "search_exhausted": result.exhausted,
        "visited_nodes": result.visited_nodes,
        "propagation_rounds": result.propagation_rounds,
        "all_solutions_valid": all(validations),
        "solutions": [[list(row) for row in matrix] for matrix in result.solutions],
    }
    (output_dir / "problem4_summary.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return payload


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/problem4"))
    args = parser.parse_args()
    payload = run(args.input, args.output_dir)
    print(
        "Problem 4:",
        f"solutions={payload['solution_count']}",
        f"unique={payload['is_unique']}",
        f"uncertain={payload['uncertain_count']}",
        f"nodes={payload['visited_nodes']}",
        f"valid={payload['all_solutions_valid']}",
    )


if __name__ == "__main__":
    main()
