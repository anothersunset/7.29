"""Enumerate and analyze the three Problem 2 worksheets."""
from __future__ import annotations
import argparse,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from src.analysis import analyze_solutions
from src.excel_io import clue_text,load_puzzles
from src.solver import solve
from src.validation import validate_matrix
SHEETS=("第二问1","第二问2","第二问3")
def run_problem2(input_path:Path,output_dir:Path):
    results=[]
    for puzzle in load_puzzles(input_path,SHEETS):
        solved=solve(puzzle.row_clues,puzzle.col_clues);validations=tuple(validate_matrix(m,puzzle.row_clues,puzzle.col_clues) for m in solved.solutions)
        if not solved.exhausted or not all(v.valid for v in validations):raise RuntimeError(f"{puzzle.name}: incomplete or invalid search")
        analysis=analyze_solutions(solved.solutions);results.append((puzzle,solved,analysis,validations))
    output_dir.mkdir(parents=True,exist_ok=True);wb=Workbook();summary=wb.active;summary.title="汇总";summary.append(["子问题","规模","解数量","唯一","不确定格数","节点数","全部验证"])
    blue=PatternFill("solid",fgColor="DCEEFF");yellow=PatternFill("solid",fgColor="FEF3C7");black=PatternFill("solid",fgColor="1F2937");white=Font(color="FFFFFF",bold=True);center=Alignment(horizontal="center",vertical="center")
    for c in summary[1]:c.fill=blue;c.font=Font(bold=True)
    for puzzle,solved,a,vals in results:
        summary.append([puzzle.name,f"{puzzle.size}×{puzzle.size}",a.solution_count,"是" if a.is_unique else "否",a.uncertain_count,solved.visited_nodes,"通过" if all(v.valid for v in vals) else "失败"])
        def add_matrix(title,matrix):
            ws=wb.create_sheet(title[:31]);ws.sheet_view.showGridLines=False;ws.cell(1,1,"行\\列")
            for j,c in enumerate(puzzle.col_clues,2):ws.cell(1,j,clue_text(c))
            for i,c in enumerate(puzzle.row_clues,2):ws.cell(i,1,clue_text(c))
            for i,row in enumerate(matrix,2):
                for j,v in enumerate(row,2):
                    cell=ws.cell(i,j,"?" if v is None else v);cell.alignment=center
                    if v is None:cell.fill=yellow
                    elif v==1:cell.fill=black;cell.font=white
            return ws
        common=add_matrix(f"{puzzle.name}公共矩阵",a.consensus_matrix);start=puzzle.size+4;common.cell(start,1,"不确定格坐标").font=Font(bold=True)
        for k,(i,j) in enumerate(a.uncertain_cells,start+1):common.cell(k,1,i);common.cell(k,2,j)
        for idx,m in enumerate(solved.solutions,1):add_matrix(f"{puzzle.name}解{idx}",m)
    wb.save(output_dir/"problem2_results.xlsx")
    payload={"problems":[{"name":p.name,"size":p.size,"solution_count":a.solution_count,"is_unique":a.is_unique,"uncertain_count":a.uncertain_count,"uncertain_cells":[list(x) for x in a.uncertain_cells],"visited_nodes":s.visited_nodes,"solutions":[[list(r) for r in m] for m in s.solutions]} for p,s,a,_ in results]}
    (output_dir/"problem2_summary.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8");return tuple(results)
def main():
    p=argparse.ArgumentParser();p.add_argument("input",type=Path);p.add_argument("--output-dir",type=Path,default=Path("outputs/problem2"));a=p.parse_args()
    for puzzle,solved,analysis,_ in run_problem2(a.input,a.output_dir):print(f"{puzzle.name}: solutions={analysis.solution_count}, unique={analysis.is_unique}, uncertain={analysis.uncertain_count}, nodes={solved.visited_nodes}")
if __name__=="__main__":main()
