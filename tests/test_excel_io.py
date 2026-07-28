import tempfile,unittest
from pathlib import Path
from openpyxl import Workbook
from src.excel_io import load_puzzle
class ExcelTests(unittest.TestCase):
    def test_load(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/"x.xlsx"; wb=Workbook(); ws=wb.active; ws.title="样例"
            for j,v in enumerate(("1,2",5,3,3,4),2): ws.cell(1,j,v)
            for i,v in enumerate((3,"2,1",4,"2,2","2,2"),2): ws.cell(i,1,v)
            wb.save(p); puzzle=load_puzzle(p,"样例")
        self.assertEqual(puzzle.size,5); self.assertEqual(puzzle.row_clues,((3,),(2,1),(4,),(2,2),(2,2))); self.assertEqual(puzzle.col_clues,((1,2),(5,),(3,),(3,),(4,)))
    def test_bad_sheet(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/"x.xlsx"; wb=Workbook(); ws=wb.active; ws.title="坏"; ws["B1"]=1; ws["C1"]=1; ws["A2"]=1; wb.save(p)
            with self.assertRaises(KeyError): load_puzzle(p,"无")
            with self.assertRaises(ValueError): load_puzzle(p,"坏")
