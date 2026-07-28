"""Excel input/output for reconstruction puzzles."""
from __future__ import annotations
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .encoding import parse_clue
from .solver import Clue, Matrix
from .validation import ValidationReport

@dataclass(frozen=True, slots=True)
class Puzzle:
    name: str; size: int; row_clues: tuple[Clue,...]; col_clues: tuple[Clue,...]
@dataclass(frozen=True, slots=True)
class Problem1Record:
    puzzle: Puzzle; matrix: Matrix; solution_count: int; exhausted: bool; visited_nodes: int; validation: ValidationReport
    @property
    def is_unique(self) -> bool: return self.solution_count == 1 and self.exhausted

def _trim(values: list[object]) -> list[object]:
    while values and values[-1] is None: values.pop()
    return values

def load_puzzle(path: str | Path, sheet_name: str) -> Puzzle:
    p=Path(path)
    if not p.is_file(): raise FileNotFoundError(p)
    wb=load_workbook(p,data_only=True,read_only=True)
    try:
        if sheet_name not in wb.sheetnames: raise KeyError(f"worksheet not found: {sheet_name}")
        ws=wb[sheet_name]
        raw_cols=_trim([ws.cell(1,c).value for c in range(2,ws.max_column+1)])
        raw_rows=_trim([ws.cell(r,1).value for r in range(2,ws.max_row+1)])
        if not raw_rows or len(raw_rows)!=len(raw_cols): raise ValueError(f"worksheet {sheet_name!r} does not define a square puzzle")
        rows=tuple(parse_clue(v) for v in raw_rows); cols=tuple(parse_clue(v) for v in raw_cols); n=len(rows)
        for axis,clues in (("row",rows),("column",cols)):
            for index,clue in enumerate(clues,1):
                if clue!=(0,) and sum(clue)+len(clue)-1>n: raise ValueError(f"{axis} clue {index} cannot fit length {n}: {clue}")
        return Puzzle(sheet_name,n,rows,cols)
    finally: wb.close()

def load_puzzles(path: str | Path,names: Sequence[str]) -> tuple[Puzzle,...]: return tuple(load_puzzle(path,n) for n in names)
def clue_text(clue: Sequence[int]) -> str: return ",".join(map(str,clue))

def write_problem1_workbook(records: Sequence[Problem1Record], output: str | Path) -> Path:
    if not records: raise ValueError("records must not be empty")
    out=Path(output); out.parent.mkdir(parents=True,exist_ok=True); wb=Workbook(); summary=wb.active; summary.title="汇总"
    headers=["子问题","规模","解数量","搜索穷尽","唯一","节点数","行验证","列验证"]; summary.append(headers)
    blue=PatternFill("solid",fgColor="DCEEFF"); green=PatternFill("solid",fgColor="DCFCE7"); red=PatternFill("solid",fgColor="FEE2E2"); black=PatternFill("solid",fgColor="1F2937")
    white=Font(color="FFFFFF",bold=True); bold=Font(bold=True); side=Side(style="thin",color="D1D5DB"); border=Border(left=side,right=side,top=side,bottom=side); center=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for c in summary[1]: c.fill=blue; c.font=bold; c.border=border; c.alignment=center
    for rec in records:
        report=rec.validation; summary.append([rec.puzzle.name,f"{rec.puzzle.size}×{rec.puzzle.size}",rec.solution_count,"是" if rec.exhausted else "否","是" if rec.is_unique else "否",rec.visited_nodes,"通过" if not report.row_mismatches else "失败","通过" if not report.col_mismatches else "失败"])
        for c in summary[summary.max_row]: c.border=border; c.alignment=center
        for col in (5,7,8): summary.cell(summary.max_row,col).fill=green if report.valid and rec.is_unique else red
        ws=wb.create_sheet(f"{rec.puzzle.name}结果"[:31]); n=rec.puzzle.size; ws.freeze_panes="B2"; ws.sheet_view.showGridLines=False; ws.cell(1,1,"行\\列")
        for j,clue in enumerate(rec.puzzle.col_clues,2): ws.cell(1,j,clue_text(clue))
        for i,clue in enumerate(rec.puzzle.row_clues,2): ws.cell(i,1,clue_text(clue))
        for i,row in enumerate(rec.matrix,2):
            for j,value in enumerate(row,2):
                c=ws.cell(i,j,value); c.border=border; c.alignment=center
                if value==1: c.fill=black; c.font=white
        for c in ws[1]: c.fill=blue; c.font=bold; c.border=border; c.alignment=center
        for i in range(2,n+2):
            c=ws.cell(i,1); c.fill=blue; c.font=bold; c.border=border; c.alignment=center
        start=n+4; ws.cell(start,1,"行验证").font=bold
        for j,label in enumerate(("序号","原始编码","重算编码","匹配"),1): ws.cell(start+1,j,label).fill=blue; ws.cell(start+1,j).font=bold
        for idx,(expected,actual) in enumerate(zip(rec.puzzle.row_clues,report.calculated_row_clues,strict=True),1):
            for j,value in enumerate((idx,clue_text(expected),clue_text(actual),"是" if expected==actual else "否"),1): c=ws.cell(start+1+idx,j,value); c.border=border; c.alignment=center
        cstart=start+n+3; ws.cell(cstart,1,"列验证").font=bold
        for j,label in enumerate(("序号","原始编码","重算编码","匹配"),1): ws.cell(cstart+1,j,label).fill=blue; ws.cell(cstart+1,j).font=bold
        for idx,(expected,actual) in enumerate(zip(rec.puzzle.col_clues,report.calculated_col_clues,strict=True),1):
            for j,value in enumerate((idx,clue_text(expected),clue_text(actual),"是" if expected==actual else "否"),1): c=ws.cell(cstart+1+idx,j,value); c.border=border; c.alignment=center
        ws.column_dimensions["A"].width=18
        for j in range(2,n+2): ws.column_dimensions[get_column_letter(j)].width=9
    summary.freeze_panes="A2"; summary.auto_filter.ref=summary.dimensions
    for j,w in enumerate((16,12,10,14,10,12,12,12),1): summary.column_dimensions[get_column_letter(j)].width=w
    wb.save(out); return out
